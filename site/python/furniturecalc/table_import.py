"""Явный табличный профиль импорта; не распознаёт произвольные клиентские таблицы."""
from copy import deepcopy
import csv
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from furniturecalc.domain.order import ValidationError, price, unique
from furniturecalc.storage import empty_state, validate_state

PRICE_COLUMNS = ['id', 'name', 'sku', 'kind', 'unit', 'price', 'length_mm', 'width_mm', 'thickness_mm', 'source', 'date', 'revision']
SPEC_COLUMNS = ['type', 'id', 'product_id', 'name', 'code', 'length_mm', 'width_mm', 'height_mm', 'quantity', 'resource_id', 'unit', 'edge_length1', 'edge_length2', 'edge_width1', 'edge_width2', 'expense_kind']
MAX_ROWS = 5000


def read_rows(data, filename, columns):
    if len(data) > 5 * 1024 * 1024:
        raise ValidationError('Файл больше 5 МБ')
    suffix = Path(filename).suffix.lower()
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        try:
            with ZipFile(BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                    raise ValidationError('Распакованный Excel превышает 30 МБ')
            book = load_workbook(BytesIO(data), read_only=True, data_only=False)
            try:
                if book.sheetnames != ['Данные']:
                    raise ValidationError('Ожидается единственный лист «Данные» из шаблона сайта')
                sheet = book['Данные']
                if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > len(columns):
                    raise ValidationError('Данные: превышен размер шаблона (5000 строк)')
                matrix = []
                for row in sheet.iter_rows():
                    values = []
                    for cell in row:
                        if cell.data_type == 'f':
                            raise ValidationError(f'Данные!{cell.coordinate}: формулы не поддерживаются; введите значение')
                        values.append(cell.value)
                    matrix.append(values)
            finally:
                book.close()
        except (BadZipFile, KeyError, OSError, ValueError) as exc:
            raise ValidationError(f'Excel: {exc}') from exc
    elif suffix == '.csv':
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise ValidationError('CSV должен быть сохранён в UTF-8') from exc
        # Шаблон использует ;, десятичная запятая остаётся частью текстового числа.
        matrix = list(csv.reader(StringIO(text), delimiter=';'))
    else:
        raise ValidationError('Поддерживаются XLSX и CSV по шаблону сайта; XLS, PDF и произвольные ТЗ пока не поддерживаются')
    if not matrix or matrix[0] != columns:
        raise ValidationError('Данные, строка 1: заголовки должны точно соответствовать шаблону: ' + '; '.join(columns))
    if len(matrix) > MAX_ROWS + 1:
        raise ValidationError('Данные: больше 5000 строк')
    rows = []
    for line, values in enumerate(matrix[1:], 2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        if len(values) > len(columns):
            raise ValidationError(f'Данные, строка {line}: лишние столбцы')
        values += [None] * (len(columns) - len(values))
        row = {key: '' if value is None else str(value).strip() for key, value in zip(columns, values)}
        row['_line'] = line
        rows.append(row)
    if not rows:
        raise ValidationError('Данные: файл не содержит строк')
    return rows


def import_catalog(data, filename):
    rows = read_rows(data, filename, PRICE_COLUMNS)
    catalog = []
    for row in rows:
        item = {key: row[key] for key in PRICE_COLUMNS}
        try:
            price(item, 'Прайс')
        except ValidationError as exc:
            raise ValidationError(f'Данные, строка {row["_line"]}: {exc}') from exc
        catalog.append(item)
    unique(catalog, 'Данные: ID прайса')
    return catalog


def import_specification(data, filename, catalog):
    rows = read_rows(data, filename, SPEC_COLUMNS)
    unique(catalog, 'Каталог')
    lookup = {item['id']: item for item in catalog}
    state = empty_state()
    state['catalog'] = deepcopy(catalog)
    state['order']['number'] = Path(filename).stem
    products = {}
    for row in rows:
        if not row['id']:
            raise ValidationError(f'Данные, строка {row["_line"]}, id: пропущен идентификатор')
        if row['type'] not in ('product', 'part', 'hardware', 'work', 'expense'):
            raise ValidationError(f'Данные, строка {row["_line"]}, type: неизвестный тип строки')
        allowed = {
            'product': {'type','id','name','code','length_mm','width_mm','height_mm','quantity'},
            'part': {'type','id','product_id','name','length_mm','width_mm','quantity','resource_id','edge_length1','edge_length2','edge_width1','edge_width2'},
            'hardware': {'type','id','product_id','quantity','resource_id','unit'},
            'work': {'type','id','product_id','quantity','resource_id','unit'},
            'expense': {'type','id','quantity','resource_id','unit','expense_kind'},
        }[row['type']]
        for field in SPEC_COLUMNS:
            if field not in allowed and row[field]:
                raise ValidationError(f'Данные, строка {row["_line"]}, {field}: поле не поддерживается для {row["type"]}; данные не будут проигнорированы')
        if row['type'] == 'product':
            if row['id'] in products:
                raise ValidationError(f'Данные, строка {row["_line"]}, id: повтор изделия {row["id"]}')
            products[row['id']] = dict(id=row['id'], name=row['name'], code=row['code'], length_mm=row['length_mm'],
                width_mm=row['width_mm'], height_mm=row['height_mm'], quantity=row['quantity'], parts=[], hardware=[], work=[], composition_confirmed=False)
    if not products:
        raise ValidationError('Данные: нет строк product (изделия)')

    def resource(row, key, kinds):
        ident = row[key]
        if not ident:
            return None
        if ident not in lookup:
            raise ValidationError(f'Данные, строка {row["_line"]}, {key}: нет точного ID «{ident}» в прайсе')
        item = lookup[ident]
        try:
            price(item, key, kinds)
        except ValidationError as exc:
            raise ValidationError(f'Данные, строка {row["_line"]}: {exc}') from exc
        return deepcopy(item)

    seen = set()
    for row in rows:
        kind = row['type']
        if kind == 'product':
            continue
        pid = row['product_id']
        key = (pid, row['id'])
        if key in seen:
            raise ValidationError(f'Данные, строка {row["_line"]}, id: повтор {row["id"]} в изделии {pid}')
        seen.add(key)
        if kind != 'expense' and pid not in products:
            raise ValidationError(f'Данные, строка {row["_line"]}, product_id: неизвестное изделие «{pid}»')
        if kind == 'part':
            products[pid]['parts'].append(dict(id=row['id'], name=row['name'], length_mm=row['length_mm'], width_mm=row['width_mm'], quantity=row['quantity'],
                material=resource(row, 'resource_id', ('sheet', 'material')), edges={side: resource(row, 'edge_' + side, ('edge',)) for side in ('length1','length2','width1','width2')}, edges_confirmed=False))
        else:
            item = dict(id=row['id'], quantity=row['quantity'], unit=row['unit'], resource=resource(row, 'resource_id', ('expense','work') if kind=='expense' else (kind,)))
            if kind == 'expense':
                if pid:
                    raise ValidationError(f'Данные, строка {row["_line"]}, product_id: расходы задаются на весь заказ, оставьте пустым')
                if row['expense_kind'] not in ('delivery','installation','other'):
                    raise ValidationError(f'Данные, строка {row["_line"]}, expense_kind: delivery / installation / other')
                item['kind'] = row['expense_kind']
                state['order']['expenses'].append(item)
            else:
                products[pid][kind].append(item)
    state['order']['products'] = list(products.values())
    return validate_state(state)
