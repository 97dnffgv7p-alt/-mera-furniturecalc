"""Собрать статический сайт; бинарные зависимости загружаются с проверкой SHA256."""
from pathlib import Path
from hashlib import sha256
from urllib.request import urlopen
import csv
import json
import shutil
import sys
import zipfile

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'site'
OUTPUT = ROOT / 'dist'
shutil.copytree(SOURCE, OUTPUT, dirs_exist_ok=True)
manifest = json.loads((SOURCE / 'python-manifest.json').read_text(encoding='utf-8'))


def download(url, relative):
    target = OUTPUT / relative
    target.parent.mkdir(parents=True, exist_ok=True)
    expected = manifest['sha256'][relative]
    if target.exists() and sha256(target.read_bytes()).hexdigest() == expected:
        return
    with urlopen(url, timeout=90) as response:
        data = response.read()
    if sha256(data).hexdigest() != expected:
        raise RuntimeError('SHA256 mismatch: ' + relative)
    target.write_bytes(data)
    print('Verified:', relative, flush=True)


for relative in manifest['sha256']:
    if relative.startswith('runtime/'):
        download('https://cdn.jsdelivr.net/pyodide/v314.0.7/full/' + relative.split('/', 1)[1], relative)
for package, version, filename in [('et_xmlfile','2.0.0','et_xmlfile-2.0.0-py3-none-any.whl'), ('openpyxl','3.1.5','openpyxl-3.1.5-py2.py3-none-any.whl')]:
    with urlopen(f'https://pypi.org/pypi/{package}/{version}/json', timeout=60) as response:
        metadata = json.load(response)
    wheel = next(item for item in metadata['urls'] if item['filename'] == filename)
    download(wheel['url'], 'wheels/' + filename)
vendor = ROOT / '.build-vendor'
vendor.mkdir(exist_ok=True)
for filename in manifest['wheels']:
    with zipfile.ZipFile(OUTPUT / 'wheels' / filename) as archive:
        archive.extractall(vendor)
sys.path.insert(0, str(vendor))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
for name in ('price', 'specification'):
    with (SOURCE / 'templates' / (name + '.csv')).open(encoding='utf-8-sig', newline='') as stream:
        rows = list(csv.reader(stream, delimiter=';'))
    book = Workbook()
    sheet = book.active
    sheet.title = 'Данные'
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='234C3A')
        cell.font = Font(color='FFFFFF', bold=True)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(45, max(15, max(len(str(cell.value or '')) for cell in column) + 2))
    book.save(OUTPUT / 'templates' / (name + '.xlsx'))
for relative, expected in manifest['sha256'].items():
    if sha256((OUTPUT / relative).read_bytes()).hexdigest() != expected:
        raise RuntimeError('Build validation failed: ' + relative)
print('FurnitureCalc static build ready', flush=True)
